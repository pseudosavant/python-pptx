from io import BytesIO

import pytest

from pptx import Presentation
from pptx.util import Inches


def roundtrip(prs):
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    return Presentation(stream)


def test_style_reference_preserves_content_and_flags():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, 0, 0, Inches(3), Inches(2)).table
    table.cell(0, 0).text = "Heading"
    table.first_row = True
    guid = "{B301B821-A1FF-4177-AEE7-76D212191A09}"
    table.style_id = guid.lower()
    with pytest.raises(ValueError):
        table.style_id = "invalid"
    reopened = roundtrip(prs).slides[0].shapes[0].table
    assert reopened.style_id == guid
    assert reopened.first_row
    assert reopened.cell(0, 0).text == "Heading"
    table.style_id = None
    assert roundtrip(prs).slides[0].shapes[0].table.style_id is None

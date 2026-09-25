from io import BytesIO

import pytest

from pptx import Presentation
from pptx.util import Inches


def roundtrip(prs):
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    return Presentation(stream)


from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.oxml.ns import qn


def test_theme_edits_preserve_other_definitions_and_hyperlink_order():
    prs = Presentation()
    master = prs.slide_master
    theme = master.theme
    original_other_color = theme.color_scheme[MSO_THEME_COLOR.ACCENT_2]
    fonts = theme.font_scheme
    theme.name = "Custom"
    theme.color_scheme.name = "Palette"
    theme.color_scheme[MSO_THEME_COLOR.ACCENT_1] = RGBColor(1, 2, 3)
    fonts.major_latin = "Arial"
    fonts.minor_latin = "Calibri"
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    run = slide.shapes.title.text_frame.paragraphs[0].add_run()
    run.text = "Theme text"
    run.hyperlink.address = "https://example.com"
    run.font.theme_font = "major"
    rpr = run._r.get_or_add_rPr()
    assert list(rpr).index(rpr.find(qn("a:cs"))) < list(rpr).index(rpr.find(qn("a:hlinkClick")))
    reopened = roundtrip(prs)
    assert reopened.slide_master.theme.name == "Custom"
    scheme = reopened.slide_master.theme.color_scheme
    assert scheme.name == "Palette"
    assert scheme[MSO_THEME_COLOR.ACCENT_1] == RGBColor(1, 2, 3)
    assert scheme[MSO_THEME_COLOR.ACCENT_2] == original_other_color
    assert reopened.slide_master.theme.font_scheme.major_latin == "Arial"
    assert reopened.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].font.theme_font == "major"
    before = master._element.xml
    assert master.text_style_font("body").size.pt > 0
    assert master._element.xml == before
    with pytest.raises(ValueError):
        run.font.theme_font = "invalid"
    assert run.font.theme_font == "major"
    run.font.theme_font = None
    assert run.font.theme_font is None

"""ScreenTips on text, shape click, and shape hover hyperlinks."""

import io

import pytest
from pptx import Presentation
from pptx.action import ActionSetting
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches


def link(prs, kind):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(2), Inches(1))
    if kind == "text":
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = "Documentation"
        return run.hyperlink
    return (
        shape.click_action
        if kind == "click"
        else ActionSetting(shape._element._nvXxPr.cNvPr, shape, hover=True)
    ).hyperlink


@pytest.mark.parametrize("kind", ["text", "click", "hover"])
def test_screen_tip_round_trip_and_clear(kind):
    prs = Presentation()
    hyperlink = link(prs, kind)
    hyperlink.address = "https://example.com"
    hyperlink.screen_tip = 'Guide <&> "notes" café'
    output = io.BytesIO()
    prs.save(output)
    shape = Presentation(output).slides[0].shapes[0]
    restored = (
        shape.text_frame.paragraphs[0].runs[0].hyperlink
        if kind == "text"
        else (
            shape.click_action
            if kind == "click"
            else ActionSetting(shape._element._nvXxPr.cNvPr, shape, hover=True)
        ).hyperlink
    )
    assert restored.screen_tip == 'Guide <&> "notes" café'
    assert restored.address == "https://example.com"
    restored.screen_tip = ""
    assert restored.screen_tip is None
    assert restored.address == "https://example.com"
    restored.screen_tip = "Again"
    restored.screen_tip = None
    assert restored.address == "https://example.com"
    assert restored.screen_tip is None


@pytest.mark.parametrize("kind", ["text", "click", "hover"])
def test_missing_link_is_not_created_by_screen_tip_access(kind):
    prs = Presentation()
    hyperlink = link(prs, kind)
    slide = prs.slides[0]
    before = slide._element.xml
    assert hyperlink.screen_tip is None
    hyperlink.screen_tip = None
    hyperlink.screen_tip = ""
    with pytest.raises(ValueError):
        hyperlink.screen_tip = "No target"
    with pytest.raises(TypeError):
        hyperlink.screen_tip = 123
    assert before == slide._element.xml


@pytest.mark.parametrize("kind", ["text", "click", "hover"])
def test_address_removal_clears_tip(kind):
    prs = Presentation()
    hyperlink = link(prs, kind)
    hyperlink.address = "https://example.com"
    hyperlink.screen_tip = "Old target"
    hyperlink.address = "https://example.org"
    assert hyperlink.screen_tip is None
    hyperlink.screen_tip = "New target"
    hyperlink.address = None
    assert hyperlink.screen_tip is None
    assert hyperlink.address is None

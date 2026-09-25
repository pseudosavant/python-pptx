from io import BytesIO

import pytest

from pptx import Presentation
from pptx.util import Inches


def roundtrip(prs):
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    return Presentation(stream)


from pptx.enum.text import MSO_NUMBERED_BULLET_STYLE
from pptx.util import BulletStyle, Pt
from pptx.oxml.xmlchemy import OxmlElement


def test_list_formatting_roundtrip_and_reset():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    p = slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text_frame.paragraphs[0]
    p.text = "Item"
    p.space_before = Pt(12)
    p.font.bold = True
    p.bullet = BulletStyle.numbered(MSO_NUMBERED_BULLET_STYLE.ARABIC_PERIOD, start_at=12)
    p.left_indent = Inches(0.5)
    p.first_line_indent = Inches(-0.2)
    p.end_font.size = Pt(18)
    reopened = roundtrip(prs).slides[0].shapes[0].text_frame.paragraphs[0]
    assert reopened.bullet == p.bullet
    assert reopened.bullet.start_at == 12
    assert reopened.left_indent == Inches(0.5)
    assert reopened.first_line_indent == Inches(-0.2)
    assert reopened.end_font.size == Pt(18)
    assert reopened.space_before == Pt(12)
    props = p._p.get_or_add_pPr()
    props.insert(0, OxmlElement("a:buFontTx"))
    p.bullet = BulletStyle.DEFAULT
    assert not props.xpath("./a:buAutoNum | ./a:buFontTx")
    assert p.font.bold
    p.bullet = BulletStyle.NO_BULLET
    p.left_indent = None
    p.first_line_indent = None
    assert p.left_indent is None and p.first_line_indent is None


@pytest.mark.parametrize("value", [0, 32768, True, 1.5])
def test_reject_invalid_numbering_start(value):
    with pytest.raises((ValueError, TypeError)):
        BulletStyle.numbered(MSO_NUMBERED_BULLET_STYLE.ARABIC_PERIOD, start_at=value)

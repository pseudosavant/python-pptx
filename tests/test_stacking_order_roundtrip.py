from io import BytesIO

import pytest

from pptx import Presentation
from pptx.util import Inches


def roundtrip(prs):
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    return Presentation(stream)


from pptx.oxml.xmlchemy import OxmlElement


@pytest.mark.parametrize("grouped", [False, True])
def test_reorder_preserves_container_and_shape_identity(grouped):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
    first = shapes.add_textbox(0, 0, Inches(1), Inches(1))
    second = shapes.add_textbox(0, 0, Inches(1), Inches(1))
    first.text = "First"
    second.text = "Second"
    tree = first.element.getparent()
    tree.append(OxmlElement("p:extLst"))
    second.send_to_back()
    second.send_to_back()
    assert [s.shape_id for s in shapes] == [second.shape_id, first.shape_id]
    assert tree[-1].tag.endswith("extLst")
    second.bring_to_front()
    second.bring_to_front()
    assert [s.shape_id for s in shapes] == [first.shape_id, second.shape_id]
    reopened = roundtrip(prs).slides[0].shapes
    if grouped:
        reopened = reopened[0].shapes
    assert [s.text for s in reopened] == ["First", "Second"]

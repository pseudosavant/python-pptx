"""Check the built distribution through its installed public API."""

from importlib.metadata import version
from io import BytesIO

import pptx
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_NUMBERED_BULLET_STYLE, MSO_TEXT_STRIKE_TYPE
from pptx.util import BulletStyle, Inches

assert version("ps-python-pptx") == pptx.__version__
prs = Presentation()
prs.slides.clear()
prs.slide_master.theme.color_scheme[MSO_THEME_COLOR.ACCENT_1] = RGBColor(1, 2, 3)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.show_master_shapes = False
slide.background.fill.set_gradient([(0, RGBColor(1, 2, 3)), (1, MSO_THEME_COLOR.ACCENT_2)])
shape = slide.shapes.add_textbox(0, 0, Inches(2), Inches(1))
shape.alt_text = "Description"
shape.alt_text_title = "Title"
shape.send_to_back()
p = shape.text_frame.paragraphs[0]
p.bullet = BulletStyle.numbered(MSO_NUMBERED_BULLET_STYLE.ARABIC_PERIOD, start_at=2)
p.left_indent = Inches(0.5)
p.first_line_indent = Inches(-0.2)
run = p.add_run()
run.text = "Wheel test"
run.font.strike = MSO_TEXT_STRIKE_TYPE.SINGLE
run.font.baseline = 0.3
run.font.theme_font = "major"
buffer = BytesIO()
prs.save(buffer)
buffer.seek(0)
reopened = Presentation(buffer)
assert reopened.slides[0].shapes[0].alt_text_title == "Title"
assert reopened.slides[0].shapes[0].text_frame.paragraphs[0].bullet.start_at == 2
print("Installed fork wheel passed")

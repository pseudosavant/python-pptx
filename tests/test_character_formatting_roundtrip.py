from io import BytesIO

import pytest

from pptx import Presentation
from pptx.util import Inches


def roundtrip(prs):
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    return Presentation(stream)


from pptx.enum.text import MSO_TEXT_STRIKE_TYPE


@pytest.mark.parametrize("baseline", [None, 0, 0.3, -0.25])
@pytest.mark.parametrize("strike", [None, *MSO_TEXT_STRIKE_TYPE])
def test_character_formatting_roundtrip(baseline, strike):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    run = slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text_frame.paragraphs[0].add_run()
    run.text = "Text"
    run.font.bold = True
    run.font.baseline = 0.5
    run.font.strike = MSO_TEXT_STRIKE_TYPE.SINGLE
    run.font.baseline = baseline
    run.font.strike = strike
    font = roundtrip(prs).slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font
    assert font.baseline == baseline
    assert font.strike == strike
    assert font.bold is True

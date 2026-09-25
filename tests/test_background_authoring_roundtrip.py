from io import BytesIO

import pytest

from pptx import Presentation
from pptx.util import Inches


def roundtrip(prs):
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    return Presentation(stream)


from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL, MSO_THEME_COLOR


@pytest.mark.parametrize("radial", [False, True])
def test_explicit_gradients_roundtrip_and_validate_atomically(radial):
    prs = Presentation()
    fill = prs.slide_master.background.fill
    stops = [(0, RGBColor(1, 2, 3)), (0.4, MSO_THEME_COLOR.ACCENT_1), (1, RGBColor(4, 5, 6))]
    fill.set_gradient(stops, angle=45, radial=radial)
    with pytest.raises(ValueError):
        fill.set_gradient(list(reversed(stops)))
    reopened = roundtrip(prs).slide_master.background.fill
    assert reopened.type == MSO_FILL.GRADIENT
    assert [s.position for s in reopened.gradient_stops] == [0, 0.4, 1]
    assert reopened.gradient_stops[0].color.rgb == RGBColor(1, 2, 3)
    assert reopened.gradient_stops[1].color.theme_color == MSO_THEME_COLOR.ACCENT_1
    if radial:
        with pytest.raises(ValueError):
            _ = reopened.gradient_angle
    else:
        assert reopened.gradient_angle == 45


def test_embedded_background_on_master_and_slide():
    prs = Presentation()
    stream = BytesIO()
    Image.new("RGB", (4, 4), "blue").save(stream, format="PNG")
    stream.seek(0)
    fill = prs.slide_master.background.fill
    prs.slide_master.set_background_picture(stream)
    assert fill.type == MSO_FILL.PICTURE
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    stream.seek(0)
    slide.set_background_picture(stream)
    reopened = roundtrip(prs)
    assert reopened.slide_master.background.fill.type == MSO_FILL.PICTURE
    assert reopened.slides[0].background.fill.type == MSO_FILL.PICTURE
    for obj in (reopened.slide_master, reopened.slides[0]):
        embed = obj._element.xpath(".//a:blip")[0].rEmbed
        assert obj.part.related_part(embed).image.size == (4, 4)

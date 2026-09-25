from io import BytesIO

from pptx import Presentation


def roundtrip(prs):
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    return Presentation(stream)


def test_clear_retains_templates_and_allows_new_slides():
    prs = Presentation()
    layout = prs.slide_layouts[0]
    first = prs.slides.add_slide(layout)
    second = prs.slides.add_slide(layout)
    first.notes_slide.notes_text_frame.text = "Notes"
    first.shapes.title.click_action.target_slide = second
    prs.slides.clear()
    assert len(prs.slides) == 0
    assert len(roundtrip(prs).slides) == 0
    slide = prs.slides.add_slide(layout)
    slide.show_master_shapes = False
    reopened = roundtrip(prs)
    assert len(reopened.slides) == 1
    assert len(reopened.slide_masters) == 1
    assert reopened.slides[0].show_master_shapes is False
    slide.show_master_shapes = None
    assert roundtrip(prs).slides[0].show_master_shapes is None

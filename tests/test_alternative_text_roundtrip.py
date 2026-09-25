from io import BytesIO

import pytest

from pptx import Presentation
from pptx.util import Inches


def roundtrip(prs):
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    return Presentation(stream)


@pytest.mark.parametrize("property_name", ["alt_text", "alt_text_title"])
def test_alternative_text_roundtrip(property_name):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(0, 0, Inches(2), Inches(1))
    name = shape.name
    assert getattr(shape, property_name) is None
    setattr(shape, property_name, 'A < B & "C"')
    reopened = roundtrip(prs).slides[0].shapes[0]
    assert getattr(reopened, property_name) == 'A < B & "C"'
    assert reopened.name == name
    with pytest.raises(TypeError):
        setattr(shape, property_name, 12)
    assert getattr(shape, property_name) == 'A < B & "C"'
    setattr(shape, property_name, "")
    assert getattr(roundtrip(prs).slides[0].shapes[0], property_name) == ""
    setattr(shape, property_name, None)
    delattr(shape, property_name)
    assert getattr(roundtrip(prs).slides[0].shapes[0], property_name) is None

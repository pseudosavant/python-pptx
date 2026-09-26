"""Round-trip checks for the read-only background snapshot API."""

import io

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.oxml.xmlchemy import OxmlElement


def test_inherited_default_background_does_not_change_xml():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    before = (slide._element.xml, slide.slide_layout._element.xml, prs.slide_master._element.xml)
    info = slide.background_info
    assert info.kind == "solid"
    assert info.color == RGBColor(255, 255, 255)
    assert before == (
        slide._element.xml,
        slide.slide_layout._element.xml,
        prs.slide_master._element.xml,
    )


def test_layout_and_slide_overrides_resolve_theme_colors():
    prs = Presentation()
    layout = prs.slide_layouts[1]
    layout.background.fill.solid()
    layout.background.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
    prs.slide_master.theme.color_scheme[MSO_THEME_COLOR.ACCENT_1] = RGBColor(20, 50, 80)
    slide = prs.slides.add_slide(layout)
    assert slide.background_info.color == RGBColor(20, 50, 80)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(40, 50, 60)
    assert slide.background_info.color == RGBColor(40, 50, 60)
    target = io.BytesIO()
    prs.save(target)
    assert Presentation(target).slides[0].background_info.color == RGBColor(40, 50, 60)


def test_gradient_snapshot_includes_geometry_and_resolved_stops():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    stops = [(0, RGBColor(0, 0, 0)), (1, RGBColor(255, 255, 255))]
    slide.background.fill.set_gradient(stops, angle=135)
    info = slide.background_info
    assert info.kind == "gradient"
    assert info.stops == tuple(stops)
    assert info.angle == 135
    assert not info.radial
    slide.background.fill.set_gradient(stops, radial=True, center=(0.25, 0.75))
    info = slide.background_info
    assert info.radial and info.center == (0.25, 0.75)


def test_picture_snapshot_returns_embedded_bytes():
    image = io.BytesIO()
    Image.new("RGB", (20, 10), "blue").save(image, format="PNG")
    prs = Presentation()
    prs.slide_master.set_background_picture(image)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    before = slide._element.xml
    info = slide.background_info
    assert info.kind == "picture"
    assert info.image == image.getvalue()
    assert before == slide._element.xml


def test_unsupported_fill_is_reported_without_mutation():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.background.fill.patterned()
    before = slide._element.xml
    assert slide.background_info.kind == "unknown"
    assert "unsupported" in slide.background_info.reason
    assert before == slide._element.xml


def test_brightness_transforms_are_resolved():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.background.fill.solid()
    color = slide.background.fill.fore_color
    color.rgb = RGBColor(255, 255, 255)
    color.brightness = -0.5
    assert slide.background_info.color == RGBColor(128, 128, 128)
    alpha = OxmlElement("a:alpha")
    alpha.set("val", "50000")
    color._color._xClr.append(alpha)
    assert slide.background_info.kind == "unknown"
